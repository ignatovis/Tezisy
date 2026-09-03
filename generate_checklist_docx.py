"""
Генерация Word-документа с чек-листами подготовки тезисов.
Два раздела: РНС и СМР.
"""

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

COMMON_BLOCKS = [
    {
        'num': '2.1',
        'title': 'Заголовок отчёта',
        'description': 'Название проекта, дата актуализации, режим (СМР / до РНС)',
        'source': 'Свойства .mpp-файла (Title, Status Date)',
    },
    {
        'num': '2.2',
        'title': 'Ключевые метрики',
        'description': 'Средний % завершения, количество критических задач, отклонение итоговой вехи',
        'source': 'Расчёт по всем рабочим задачам .mpp',
    },
    {
        'num': '2.3',
        'title': 'Ближайшие события (горизонт 1 месяц)',
        'description': 'Вехи и задачи, завершающиеся в ближайший месяц',
        'source': 'Фильтрация задач по дате финиша',
    },
]

SMR_BLOCKS = [
    {
        'num': '3.1',
        'title': 'Декомпозиция отклонений по корпусам',
        'description': 'Список задач с отклонением от базового плана более 14 дней, сгруппированных по корпусам/зданиям',
        'what_shows': 'Название задачи, корпус, величина отклонения в днях, % завершения',
    },
    {
        'num': '3.2',
        'title': 'Отклонения по зонам / этапам работ',
        'description': 'Агрегация отклонений по зонам (подготовительный период, фундамент, каркас, фасад, кровля, инженерия, отделка, благоустройство)',
        'what_shows': 'Зона, количество задач с отклонениями, среднее отклонение, максимальное отклонение',
    },
    {
        'num': '3.3',
        'title': 'Критические задачи с низким % завершения',
        'description': 'Задачи на критическом пути с % завершения ниже 50%, требующие немедленного внимания',
        'what_shows': 'Название задачи, % завершения, отклонение, крайний срок',
    },
    {
        'num': '3.4',
        'title': 'Временной резерв по некритическим задачам',
        'description': 'Некритические задачи с положительным total slack — потенциал для переброски ресурсов',
        'what_shows': 'Название задачи, временной резерв (дни), текущее отклонение',
    },
    {
        'num': '3.5',
        'title': 'Зоны опережения и временного резерва',
        'description': 'Задачи с опережением базового плана более 14 дней — возможность перераспределить ресурсы',
        'what_shows': 'Название задачи, величина опережения (дни)',
    },
    {
        'num': '3.6',
        'title': 'Влияние на финальные вехи (ЗОС / РНВ)',
        'description': 'Прогноз дат получения ЗОС и РНВ, отклонение от базового плана',
        'what_shows': 'Веха, прогнозная дата, отклонение',
    },
]

RNS_BLOCKS = [
    {
        'num': '4.1',
        'title': 'Критический путь до РНС',
        'description': 'Критические вехи на пути к получению разрешения на строительство, их статус и отклонения',
        'what_shows': 'Веха, прогнозная дата, базовая дата, отклонение, крайний срок',
    },
    {
        'num': '4.2',
        'title': 'Анализ причин сдвигов',
        'description': 'Корневые причины отклонений от базового плана: какие задачи сдвинулись и на сколько',
        'what_shows': 'Название задачи, отклонение (дни), % завершения, корневая причина',
    },
    {
        'num': '4.3',
        'title': 'Каскадное влияние сдвигов',
        'description': 'Цепочка предшественников вехи РНС — как задержки на ранних этапах каскадно сдвигают дату РНС',
        'what_shows': 'Цепочка задач, отклонение каждого звена, суммарное влияние на РНС',
    },
    {
        'num': '4.4',
        'title': 'Обязательства по МПТ',
        'description': 'Задачи, связанные с межведомственными и межпроектными требованиями (МПТ), их статус',
        'what_shows': 'Задача МПТ, дата завершения, отклонение, % готовности',
    },
    {
        'num': '4.5',
        'title': 'Зоны неопределённости',
        'description': 'Задачи без базового плана или с нулевым прогрессом при приближающейся дате — зоны риска',
        'what_shows': 'Название задачи, дата финиша, статус, причина неопределённости',
    },
    {
        'num': '4.6',
        'title': 'Прогноз даты РНС',
        'description': 'Итоговый прогноз даты получения РНС, сравнение с базовым планом и крайним сроком',
        'what_shows': 'Прогнозная дата, базовая дата, отклонение, крайний срок, необходимость компенсационных мер',
    },
]

DECISIONS_BLOCK = {
    'num': '5.1',
    'title': 'Управленческие решения',
    'description': 'Автоматически сформированные рекомендации: переброска ресурсов, запросы на согласование, компенсационные мероприятия',
    'what_shows': 'Тип решения, описание, приоритет',
}


def _set_cell_shading(cell, color_hex):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color_hex)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)


def _add_section_title(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)


def _add_checklist_table(doc, blocks, include_what_shows=True):
    cols = 4 if include_what_shows else 3
    table = doc.add_table(rows=1, cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ['№', 'Раздел', 'Описание']
    if include_what_shows:
        headers.append('Что отображается в отчёте')

    hdr_row = table.rows[0]
    for i, h in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.text = h
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _set_cell_shading(cell, '1A365D')

    for block in blocks:
        row = table.add_row()
        row.cells[0].text = block['num']
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[1].text = block['title']
        row.cells[2].text = block.get('description', '')
        if include_what_shows and 'what_shows' in block:
            row.cells[3].text = block['what_shows']
        elif include_what_shows:
            row.cells[3].text = block.get('source', '')

        for cell in row.cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)

    widths = [Cm(1.2), Cm(4.5), Cm(6.5)]
    if include_what_shows:
        widths.append(Cm(5.5))
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = w


def generate():
    doc = Document()

    style = doc.styles['Normal']
    style.font.name = 'Segoe UI'
    style.font.size = Pt(10)
    style.paragraph_format.space_after = Pt(4)

    title = doc.add_heading('Чек-листы подготовки тезисов', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    doc.add_paragraph(
        'Описание разделов автоматического отчёта по план-графику (.mpp). '
        'Каждый раздел формируется макросом автоматически на основе анализа задач план-графика.'
    )

    doc.add_paragraph()

    # === Общие блоки ===
    _add_section_title(doc, 'Общие разделы (оба режима)', level=1)
    doc.add_paragraph('Разделы, включаемые в отчёт независимо от выбранного режима.')
    _add_checklist_table(doc, COMMON_BLOCKS, include_what_shows=True)

    doc.add_paragraph()

    # === СМР ===
    _add_section_title(doc, 'Режим СМР — Строительно-монтажные работы', level=1)
    doc.add_paragraph(
        'Режим для проектов на стадии строительства. '
        'Фокус: отклонения от графика, критический путь, ресурсы, влияние на ЗОС/РНВ.'
    )
    _add_checklist_table(doc, SMR_BLOCKS, include_what_shows=True)

    doc.add_paragraph()

    # === РНС ===
    _add_section_title(doc, 'Режим РНС — До получения разрешения на строительство', level=1)
    doc.add_paragraph(
        'Режим для проектов на стадии получения разрешительной документации. '
        'Фокус: критический путь до РНС, каскадные сдвиги, обязательства по МПТ, прогноз даты РНС.'
    )
    _add_checklist_table(doc, RNS_BLOCKS, include_what_shows=True)

    doc.add_paragraph()

    # === Управленческие решения ===
    _add_section_title(doc, 'Управленческие решения (оба режима)', level=1)
    doc.add_paragraph(
        'Формируются автоматически на основе выявленных отклонений и рисков.'
    )
    _add_checklist_table(doc, [DECISIONS_BLOCK], include_what_shows=True)

    out_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                            'Чек-листы_подготовки_тезисов.docx')
    doc.save(out_path)
    return out_path


if __name__ == '__main__':
    path = generate()
    print(f'Готово: {path}')
