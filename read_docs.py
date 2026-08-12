import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

from docx import Document
import os

base = os.path.dirname(os.path.abspath(__file__))

for fname in ["Текущий макрос тезисов.docx", "Текущее описание.docx"]:
    path = os.path.join(base, fname)
    print(f"\n{'='*60}")
    print(f"ФАЙЛ: {fname}")
    print('='*60)

    d = Document(path)

    print("\n--- ПАРАГРАФЫ ---")
    for i, p in enumerate(d.paragraphs):
        if p.text.strip():
            print(f"[{i}] {p.text}")

    if d.tables:
        print(f"\n--- ТАБЛИЦЫ ({len(d.tables)} шт) ---")
        for ti, t in enumerate(d.tables):
            print(f"\nТаблица {ti+1}:")
            for ri, r in enumerate(t.rows):
                cells = [c.text.strip().replace('\n', ' | ') for c in r.cells]
                print(f"  R{ri}: {' || '.join(cells)}")
