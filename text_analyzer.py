"""
Анализ текста тезисов коллег по чек-листу.
Проверяет покрытие обязательных разделов и полноту описания.
"""

import re
from collections import OrderedDict


COMMON_SECTIONS = [
    {
        'num': '2.1',
        'title': 'Заголовок отчёта',
        'description': 'Название проекта, дата актуализации, режим (СМР / до РНС)',
        'patterns': [
            r'проект[а-яё]*',
            r'актуализаци[а-яё]*',
            r'\d{2}[.\-/]\d{2}[.\-/]\d{4}',
            r'режим',
        ],
        'min_matches': 2,
        'weight': 1.0,
    },
    {
        'num': '2.2',
        'title': 'Ключевые метрики',
        'description': 'Средний % завершения, количество критических задач, отклонение итоговой вехи',
        'patterns': [
            r'процент|%|завершени[а-яё]*',
            r'критическ[а-яё]*\s+задач',
            r'отклонени[а-яё]*',
            r'вех[а-яё]*',
            r'\d+\s*%',
        ],
        'min_matches': 2,
        'weight': 1.5,
    },
    {
        'num': '2.3',
        'title': 'Ближайшие события (горизонт 1 месяц)',
        'description': 'Вехи и задачи, завершающиеся/начинающиеся в ближайший месяц',
        'patterns': [
            r'событи[а-яё]*',
            r'завершен[а-яё]*|завершающ[а-яё]*',
            r'начал[а-яё]*|начинающ[а-яё]*',
            r'месяц|ближайш[а-яё]*',
            r'горизонт',
        ],
        'min_matches': 2,
        'weight': 1.0,
    },
]

SMR_SECTIONS = [
    {
        'num': '3.1',
        'title': 'Декомпозиция отклонений',
        'description': 'Задачи с отклонением >14 дней, группировка по корпусам/направлениям',
        'patterns': [
            r'декомпозици[а-яё]*|отклонени[а-яё]*',
            r'корпус[а-яё]*|направлени[а-яё]*',
            r'задержк[а-яё]*|задержа[а-яё]*',
            r'\d+\s*дн[а-яё]*',
            r'максимальн[а-яё]*\s+отклонени',
        ],
        'min_matches': 2,
        'weight': 2.0,
    },
    {
        'num': '3.2',
        'title': 'Группировка по корпусам/зонам',
        'description': 'Агрегация отклонений по корпусам: среднее, максимальное отклонение, % завершения',
        'patterns': [
            r'корпус[а-яё]*|зон[а-яё]*',
            r'группировк[а-яё]*|агрегаци[а-яё]*',
            r'средн[а-яё]*\s+отклонени',
            r'отстающ[а-яё]*|опережающ[а-яё]*',
        ],
        'min_matches': 2,
        'weight': 1.5,
    },
    {
        'num': '3.3',
        'title': 'Критические задачи с низким % завершения',
        'description': 'Задачи на критическом пути с % завершения ниже 50%',
        'patterns': [
            r'критическ[а-яё]*\s+(задач|пут)',
            r'низк[а-яё]*\s+(%|процент)',
            r'завершен[а-яё]*\s*\d+\s*%',
            r'контрол[а-яё]*|вниман[а-яё]*',
        ],
        'min_matches': 2,
        'weight': 2.0,
    },
    {
        'num': '3.4',
        'title': 'Временной резерв',
        'description': 'Некритические задачи с положительным временным резервом',
        'patterns': [
            r'временн[а-яё]*\s+резерв',
            r'некритическ[а-яё]*',
            r'slack|запас',
            r'переброск[а-яё]*\s+ресурс',
        ],
        'min_matches': 1,
        'weight': 1.0,
    },
    {
        'num': '3.5',
        'title': 'Зоны опережения',
        'description': 'Задачи с опережением базового плана более 14 дней',
        'patterns': [
            r'опережени[а-яё]*',
            r'резерв[а-яё]*',
            r'ресурсн[а-яё]*\s+резерв',
            r'перераспредел[а-яё]*',
        ],
        'min_matches': 1,
        'weight': 1.0,
    },
    {
        'num': '3.6',
        'title': 'Влияние на финальные вехи (ЗОС/РНВ)',
        'description': 'Прогноз дат получения ЗОС и РНВ',
        'patterns': [
            r'ЗОС|РНВ',
            r'финальн[а-яё]*\s+вех',
            r'прогноз[а-яё]*',
            r'влияни[а-яё]*',
        ],
        'min_matches': 2,
        'weight': 2.0,
    },
]

RNS_SECTIONS = [
    {
        'num': '4.1',
        'title': 'Критический путь до РНС',
        'description': 'Критические вехи на пути к получению разрешения на строительство',
        'patterns': [
            r'критическ[а-яё]*\s+пут',
            r'РНС',
            r'разрешени[а-яё]*\s+(на\s+)?строительств',
            r'вех[а-яё]*',
        ],
        'min_matches': 2,
        'weight': 2.0,
    },
    {
        'num': '4.2',
        'title': 'Анализ причин сдвигов',
        'description': 'Корневые причины отклонений от базового плана',
        'patterns': [
            r'причин[а-яё]*\s+сдвиг',
            r'корнев[а-яё]*\s+причин',
            r'задержк[а-яё]*|сдвиг[а-яё]*',
            r'базов[а-яё]*\s+план',
        ],
        'min_matches': 2,
        'weight': 1.5,
    },
    {
        'num': '4.3',
        'title': 'Каскадное влияние сдвигов',
        'description': 'Цепочка предшественников — как задержки каскадно сдвигают дату РНС',
        'patterns': [
            r'каскад[а-яё]*',
            r'цепочк[а-яё]*|предшественник',
            r'влияни[а-яё]*\s+(на\s+)?РНС',
            r'смещени[а-яё]*',
        ],
        'min_matches': 1,
        'weight': 1.5,
    },
    {
        'num': '4.4',
        'title': 'Обязательства по МПТ / ковенантам',
        'description': 'Задачи МПТ, их статус и риски',
        'patterns': [
            r'МПТ',
            r'ковенант[а-яё]*',
            r'обязательств[а-яё]*',
            r'межведомствен[а-яё]*|межпроектн[а-яё]*',
        ],
        'min_matches': 1,
        'weight': 1.5,
    },
    {
        'num': '4.5',
        'title': 'Зоны неопределённости',
        'description': 'Задачи без базового плана — зоны риска',
        'patterns': [
            r'неопределённост[а-яё]*|неопределенност[а-яё]*',
            r'без\s+базов[а-яё]*\s+план',
            r'риск[а-яё]*',
            r'baseline',
        ],
        'min_matches': 1,
        'weight': 1.0,
    },
    {
        'num': '4.6',
        'title': 'Прогноз даты РНС',
        'description': 'Итоговый прогноз даты получения РНС',
        'patterns': [
            r'прогноз[а-яё]*\s+(дат[а-яё]*\s+)?РНС',
            r'дат[а-яё]*\s+РНС',
            r'прогнозн[а-яё]*\s+дат',
            r'компенсационн[а-яё]*\s+мероприят',
        ],
        'min_matches': 1,
        'weight': 2.0,
    },
]

DECISIONS_SECTION = {
    'num': '5.1',
    'title': 'Управленческие решения',
    'description': 'Рекомендации: переброска ресурсов, согласования, компенсационные мероприятия',
    'patterns': [
        r'управленческ[а-яё]*\s+решени',
        r'рекомендаци[а-яё]*|предложени[а-яё]*',
        r'переброск[а-яё]*\s+ресурс',
        r'согласовани[а-яё]*',
        r'компенсационн[а-яё]*\s+мероприят',
        r'мероприяти[а-яё]*',
    ],
    'min_matches': 1,
    'weight': 1.5,
}


def detect_mode(text):
    text_lower = text.lower()
    smr_score = 0
    rns_score = 0

    smr_keywords = ['смр', 'строительно-монтажн', 'зос', 'рнв', 'фасад', 'монолит',
                     'кровля', 'отделка', 'благоустройств']
    rns_keywords = ['рнс', 'разрешение на строительство', 'гпзу', 'ппт',
                     'разрешительн', 'документаци', 'мпт', 'ковенант']

    for kw in smr_keywords:
        if kw in text_lower:
            smr_score += 1
    for kw in rns_keywords:
        if kw in text_lower:
            rns_score += 1

    if smr_score > rns_score:
        return 'smr'
    elif rns_score > smr_score:
        return 'rns'
    return 'auto'


def _count_pattern_matches(text, patterns):
    matches = 0
    matched_patterns = []
    for pattern in patterns:
        if re.search(pattern, text, re.IGNORECASE):
            matches += 1
            matched_patterns.append(pattern)
    return matches, matched_patterns


def _has_numeric_data(text):
    dates = len(re.findall(r'\d{2}[.\-/]\d{2}[.\-/]\d{4}', text))
    percentages = len(re.findall(r'\d+\s*%', text))
    days = len(re.findall(r'\d+\s*дн[а-яё]*', text))
    return {
        'dates': dates,
        'percentages': percentages,
        'days_mentions': days,
        'total': dates + percentages + days,
    }


def _score_section(text, section):
    matches, matched_patterns = _count_pattern_matches(text, section['patterns'])
    total_patterns = len(section['patterns'])
    min_matches = section.get('min_matches', 1)

    if matches == 0:
        return 0.0, 'not_found', []

    keyword_score = min(matches / total_patterns, 1.0) * 70

    paragraphs = [p.strip() for p in text.split('\n') if p.strip()]
    relevant_paras = []
    for p in paragraphs:
        for pattern in section['patterns']:
            if re.search(pattern, p, re.IGNORECASE):
                relevant_paras.append(p)
                break

    detail_score = 0
    if relevant_paras:
        avg_len = sum(len(p) for p in relevant_paras) / len(relevant_paras)
        if avg_len > 100:
            detail_score = 30
        elif avg_len > 50:
            detail_score = 20
        elif avg_len > 20:
            detail_score = 10
        else:
            detail_score = 5

    total_score = min(keyword_score + detail_score, 100)

    if matches >= min_matches:
        if total_score >= 70:
            status = 'good'
        elif total_score >= 40:
            status = 'partial'
        else:
            status = 'weak'
    else:
        status = 'weak'

    return round(total_score, 1), status, matched_patterns


def analyze_text(text, mode=None):
    if not text or not text.strip():
        return {
            'error': 'Текст для анализа не предоставлен',
            'sections': [],
            'overall_score': 0,
            'mode': None,
        }

    if mode not in ('smr', 'rns'):
        mode = detect_mode(text)
        if mode == 'auto':
            mode = 'smr'

    if mode == 'smr':
        mode_sections = SMR_SECTIONS
        mode_label = 'СМР'
    else:
        mode_sections = RNS_SECTIONS
        mode_label = 'до РНС'

    all_sections = COMMON_SECTIONS + mode_sections + [DECISIONS_SECTION]

    results = []
    total_weighted_score = 0
    total_weight = 0

    for section in all_sections:
        score, status, matched = _score_section(text, section)
        weight = section.get('weight', 1.0)
        total_weighted_score += score * weight
        total_weight += weight

        results.append({
            'num': section['num'],
            'title': section['title'],
            'description': section['description'],
            'score': score,
            'status': status,
            'weight': weight,
            'matched_count': len(matched),
            'total_patterns': len(section['patterns']),
        })

    overall_score = round(total_weighted_score / total_weight, 1) if total_weight > 0 else 0

    good_count = sum(1 for r in results if r['status'] == 'good')
    partial_count = sum(1 for r in results if r['status'] == 'partial')
    weak_count = sum(1 for r in results if r['status'] == 'weak')
    missing_count = sum(1 for r in results if r['status'] == 'not_found')

    numeric = _has_numeric_data(text)
    word_count = len(text.split())
    para_count = len([p for p in text.split('\n') if p.strip()])

    recommendations = _generate_recommendations(results, mode)

    if overall_score >= 80:
        verdict = 'Отлично'
        verdict_class = 'good'
    elif overall_score >= 60:
        verdict = 'Хорошо'
        verdict_class = 'partial'
    elif overall_score >= 40:
        verdict = 'Удовлетворительно'
        verdict_class = 'weak'
    else:
        verdict = 'Требуется доработка'
        verdict_class = 'bad'

    return {
        'mode': mode,
        'mode_label': mode_label,
        'sections': results,
        'overall_score': overall_score,
        'verdict': verdict,
        'verdict_class': verdict_class,
        'stats': {
            'good': good_count,
            'partial': partial_count,
            'weak': weak_count,
            'missing': missing_count,
            'total': len(results),
        },
        'text_stats': {
            'word_count': word_count,
            'para_count': para_count,
            'numeric': numeric,
        },
        'recommendations': recommendations,
    }


def _generate_recommendations(results, mode):
    recs = []

    missing = [r for r in results if r['status'] == 'not_found']
    if missing:
        names = ', '.join(f"«{r['num']} {r['title']}»" for r in missing)
        recs.append({
            'type': 'critical',
            'text': f'Отсутствуют обязательные разделы: {names}',
        })

    weak = [r for r in results if r['status'] == 'weak']
    if weak:
        for r in weak:
            recs.append({
                'type': 'warning',
                'text': f'Раздел «{r["num"]} {r["title"]}» раскрыт недостаточно — '
                        f'добавьте конкретные данные: цифры, даты, названия задач',
            })

    partial = [r for r in results if r['status'] == 'partial']
    if partial:
        for r in partial:
            recs.append({
                'type': 'info',
                'text': f'Раздел «{r["num"]} {r["title"]}» можно усилить — '
                        f'добавьте количественные показатели и конкретику',
            })

    has_numeric = any(r for r in results if r['score'] > 0)
    if has_numeric and not any(r['status'] == 'not_found' for r in results):
        if all(r['status'] in ('good', 'partial') for r in results):
            recs.append({
                'type': 'success',
                'text': 'Все обязательные разделы присутствуют. '
                        'Рекомендуется проверить актуальность данных.',
            })

    if not recs:
        recs.append({
            'type': 'info',
            'text': 'Проверьте соответствие текста актуальным данным из план-графика.',
        })

    return recs


def read_docx(file_path):
    from docx import Document
    doc = Document(file_path)
    paragraphs = []
    for p in doc.paragraphs:
        if p.text.strip():
            paragraphs.append(p.text.strip())

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(' | '.join(cells))

    return '\n'.join(paragraphs)
